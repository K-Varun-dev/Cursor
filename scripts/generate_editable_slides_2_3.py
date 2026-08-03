#!/usr/bin/env python3
"""Recreate PDF slides 2 and 3 as fully editable PowerPoint slides."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "editable_slides_2_3.pptx"

W, H = 13.333, 7.5
NAVY = RGBColor(0x17, 0x32, 0x4D)
BLUE = RGBColor(0x00, 0x78, 0xB8)
GREEN = RGBColor(0x2E, 0x86, 0x75)
ORANGE = RGBColor(0xE8, 0x7E, 0x22)
RED = RGBColor(0xC9, 0x4B, 0x4B)
MUTED = RGBColor(0x5E, 0x72, 0x84)
LIGHT_BORDER = RGBColor(0xC8, 0xD5, 0xDE)
BG = RGBColor(0xF5, 0xF8, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE_BLUE = RGBColor(0xED, 0xF7, 0xFC)
PALE_GREEN = RGBColor(0xEC, 0xF6, 0xF3)
PALE_ORANGE = RGBColor(0xFE, 0xF4, 0xE8)
PALE_RED = RGBColor(0xFD, 0xEC, 0xED)
CHEVRON = RGBColor(0xA9, 0xBC, 0xC9)
FONT = "Arial"


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=12,
    color=NAVY,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    italic=False,
    tracking=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    if tracking is not None:
        p.font._element.set("spc", str(tracking))
    return box


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    add_rect(slide, 0, 0, W, 0.18, BLUE)
    add_rect(slide, 12.93, 0, 0.40, H, BLUE)


def add_header(slide, title, subtitle):
    add_text(slide, 0.60, 0.46, 11.9, 0.43, title, 25, NAVY, True)
    add_text(slide, 0.60, 1.00, 11.8, 0.25, subtitle, 11, MUTED)


def add_footer(slide, source_text, number):
    add_text(slide, 0.60, 6.90, 6.0, 0.18, source_text, 6, MUTED, italic=True)
    add_text(
        slide,
        0.52,
        7.17,
        3.2,
        0.14,
        "S E N S O R Y - F R I E N D L Y   U R B A N   F U T U R E S",
        5.5,
        MUTED,
        True,
    )
    add_text(
        slide,
        10.95,
        7.17,
        1.75,
        0.14,
        f"D A T A   S C I E N C E   P R O P O S A L   {number}",
        5.5,
        MUTED,
        True,
        PP_ALIGN.RIGHT,
    )


def add_circle(slide, x, y, diameter, text, fill, size=9):
    circle = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(diameter),
        Inches(diameter),
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = fill
    circle.line.fill.background()
    tf = circle.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = WHITE


def add_card(slide, x, y, w, h, title, title_fill, bullets):
    add_rect(slide, x, y, w, h, WHITE, LIGHT_BORDER, radius=True)
    add_rect(slide, x, y, w, 0.47, title_fill)
    add_text(
        slide,
        x + 0.18,
        y + 0.10,
        w - 0.36,
        0.27,
        title,
        13,
        WHITE,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    start_y = y + 0.70
    gap = (h - 0.86) / len(bullets)
    for index, bullet in enumerate(bullets):
        by = start_y + index * gap
        add_circle(slide, x + 0.20, by, 0.14, "", title_fill)
        add_text(slide, x + 0.47, by - 0.01, w - 0.68, 0.31, bullet, 9.4, NAVY)


def slide_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Data Pipeline: Raw Counts to User-Facing Decisions",
        "The data layer must preserve provenance, quality status, freshness and uncertainty.",
    )

    steps = [
        ("Official APIs\n+ OSM snapshot", BLUE),
        ("Immutable raw\nsnapshot +\nmetadata", BLUE),
        ("Validate,\ndeduplicate\nand quality-flag", BLUE),
        ("Normalised\nPostgreSQL", GREEN),
        ("Baseline,\nforecast\nand route\nweighting", GREEN),
        ("Application API\nmap + alerts", ORANGE),
    ]
    xs = [0.63, 2.61, 4.59, 6.55, 8.53, 10.50]
    box_w = 1.64
    for i, ((label, colour), x) in enumerate(zip(steps, xs), start=1):
        fill = PALE_BLUE if colour == BLUE else PALE_GREEN if colour == GREEN else PALE_ORANGE
        add_rect(slide, x, 1.54, box_w, 1.05, fill, colour, radius=True)
        add_circle(slide, x + 0.08, 1.65, 0.29, str(i), colour, 8)
        add_text(
            slide,
            x + 0.26,
            1.84,
            box_w - 0.38,
            0.57,
            label,
            8.8,
            NAVY,
            True,
            PP_ALIGN.CENTER,
            MSO_ANCHOR.MIDDLE,
        )
        if i < len(steps):
            chev = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.CHEVRON,
                Inches(x + box_w + 0.03),
                Inches(1.86),
                Inches(0.24),
                Inches(0.37),
            )
            chev.fill.solid()
            chev.fill.fore_color.rgb = CHEVRON
            chev.line.fill.background()

    add_card(
        slide,
        0.66,
        3.03,
        5.42,
        2.75,
        "Ingestion controls",
        BLUE,
        [
            "Record source URL, retrieval time and licence",
            "Validate IDs, timestamps, coordinates and non-negative counts",
            "Upsert on (sensor_id, observed_at)",
            "Quarantine invalid records—never silently discard",
        ],
    )
    add_card(
        slide,
        6.66,
        3.03,
        5.42,
        2.75,
        "Interpretation controls",
        GREEN,
        [
            "Distinguish zero movement from sensor/API failure",
            "Preserve sensor relocation history and active periods",
            "Display freshness, geographic coverage and confidence",
            "Version raw snapshots, transformations and model outputs",
        ],
    )
    add_footer(
        slide,
        "Source basis: Open Data.pdf — planned pipeline and quality controls",
        2,
    )


def add_label_pill(slide, x, y, w, text, fill):
    add_rect(slide, x, y, w, 0.36, fill, None, radius=True)
    add_text(
        slide,
        x,
        y,
        w,
        0.36,
        text,
        8,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )


def add_level_row(slide, x, y, w, label, value, line, fill):
    add_rect(slide, x, y, w, 0.56, fill, line, radius=True)
    add_text(
        slide,
        x + 0.18,
        y + 0.16,
        1.25,
        0.20,
        label,
        10,
        line,
        True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        x + 1.55,
        y + 0.16,
        w - 1.75,
        0.20,
        value,
        9.5,
        NAVY,
        valign=MSO_ANCHOR.MIDDLE,
    )


def slide_analytics(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Explainable Analytics for the Onboarding Iteration",
        "Start with transparent baselines before introducing complex machine-learning models.",
    )

    # Two editable white cards.
    add_rect(slide, 0.66, 1.39, 5.42, 4.80, WHITE, LIGHT_BORDER, radius=True)
    add_rect(slide, 6.39, 1.39, 5.98, 4.80, WHITE, LIGHT_BORDER, radius=True)

    add_label_pill(slide, 0.90, 1.66, 2.10, "CROWD CLASSIFICATION", BLUE)
    add_text(
        slide,
        0.94,
        2.30,
        4.85,
        0.48,
        "Compare each observation with the same sensor’s\nhistorical distribution for the same weekday and hour.",
        11.3,
        NAVY,
    )

    add_level_row(slide, 0.99, 2.99, 4.68, "LOW", "≤ 40th percentile", GREEN, PALE_GREEN)
    add_level_row(
        slide,
        0.99,
        3.72,
        4.68,
        "MODERATE",
        "40th–75th percentile",
        ORANGE,
        PALE_ORANGE,
    )
    add_level_row(slide, 0.99, 4.45, 4.68, "HIGH", "> 75th percentile", RED, PALE_RED)

    add_text(slide, 0.96, 5.46, 1.42, 0.22, "Why this is better", 9.5, BLUE, True)
    add_text(
        slide,
        2.36,
        5.35,
        3.35,
        0.50,
        "It avoids one arbitrary threshold across\nlocations with very different normal pedestrian\nvolumes.",
        9,
        MUTED,
    )

    add_label_pill(slide, 6.65, 1.66, 2.15, "NEXT-HOUR FORECAST", GREEN)
    forecast = [
        "Historical median for sensor, weekday and hour",
        "Adjustment using the most recent observation",
        "Confidence from sample size and data freshness",
        "Validation using MAE and high-crowd precision/recall",
    ]
    for index, text in enumerate(forecast, start=1):
        y = 2.34 + (index - 1) * 0.68
        add_circle(slide, 6.78, y, 0.40, str(index), GREEN, 10)
        add_text(
            slide,
            7.35,
            y + 0.07,
            4.55,
            0.26,
            text,
            9.5,
            NAVY,
            valign=MSO_ANCHOR.MIDDLE,
        )

    add_rect(slide, 6.78, 5.24, 5.06, 0.57, PALE_BLUE, BLUE, radius=True)
    add_text(
        slide,
        6.95,
        5.33,
        4.72,
        0.38,
        "User output: current level + next-hour estimate + confidence +\nearly rerouting prompt",
        9,
        NAVY,
        True,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )

    add_footer(
        slide,
        "Source basis: Open Data.pdf — selected analytical approach",
        3,
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide_pipeline(prs)
    slide_analytics(prs)
    prs.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
