#!/usr/bin/env python3
"""Generate open data onboarding slides as PowerPoint (.pptx)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Theme colours (match LaTeX deck)
NAVY = RGBColor(0x17, 0x25, 0x54)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
SLATE = RGBColor(0x33, 0x41, 0x55)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
CYAN = RGBColor(0x08, 0x91, 0xB2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF8, 0xFA, 0xFC)
LIGHT_BLUE = RGBColor(0xEE, 0xF2, 0xFF)
LIGHT_CYAN = RGBColor(0xEC, 0xFE, 0xFF)

OUT = Path(__file__).resolve().parents[1] / "open_data_presentation.pptx"


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title: str) -> None:
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.85)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.35)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_text_box(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    font_size=12,
    bold=False,
    color=SLATE,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align
        p.space_after = Pt(4)
    return box


def add_bullet_box(
    slide, left, top, width, height, heading: str, bullets: list[str], body_size=11
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_BLUE
    shape.line.color.rgb = INDIGO
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.08)
    h = tf.paragraphs[0]
    h.text = heading
    h.font.size = Pt(13)
    h.font.bold = True
    h.font.color.rgb = INDIGO
    h.space_after = Pt(6)
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(body_size)
        p.font.color.rgb = SLATE
        p.space_after = Pt(3)


def slide_datasets(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GREY)
    add_title_bar(slide, "The five datasets we actually need")

    rows = [
        (
            "Past-hour pedestrian counts (per minute)",
            "API/CSV; ~15 min",
            "Current volume & high-crowd alerts",
            "Every 15 min",
            "CC BY 4.0",
        ),
        (
            "Pedestrian counts per hour (2009–present)",
            "API/CSV; monthly",
            "Baselines, quiet patterns, forecasts",
            "Monthly",
            "CC BY 4.0",
        ),
        (
            "Pedestrian sensor locations",
            "API/CSV; as maintained",
            "Map position, status, relocation notes",
            "Daily check",
            "CC BY 4.0",
        ),
        (
            "Landmarks & places of interest",
            "API/CSV; periodic",
            "Candidate parks, libraries, public breaks",
            "Per iteration",
            "CC BY 4.0",
        ),
        (
            "OpenStreetMap pedestrian network",
            "OSM extract/API",
            "Walkable graph for sensory-weighted routes",
            "Versioned snapshot",
            "ODbL 1.0",
        ),
    ]

    headers = [
        "Open dataset",
        "Source refresh",
        "Why we need it",
        "Our refresh",
        "Licence",
    ]
    tbl = slide.shapes.add_table(
        len(rows) + 1, 5, Inches(0.35), Inches(1.05), Inches(12.6), Inches(4.85)
    ).table
    col_widths = [2.35, 1.55, 3.55, 1.35, 1.1]
    for ci, w in enumerate(col_widths):
        tbl.columns[ci].width = Inches(w)

    for ci, h in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(9)
            p.font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = INDIGO

    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(8)
                p.font.color.rgb = SLATE

    add_text_box(
        slide,
        Inches(0.35),
        Inches(6.05),
        Inches(12.5),
        Inches(0.5),
        "Selection rule: every source must support a user story—not included just because it is available.",
        font_size=11,
        bold=True,
        color=NAVY,
    )


def slide_why_chosen(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GREY)
    add_title_bar(slide, "Why we chose these datasets")

    items = [
        (
            "Live minute counts",
            "Epic 1 & 2: real-time crowd and alerts. Only city open feed refreshed ~every 15 min.",
        ),
        (
            "Historical hourly counts",
            "Defines “unusually busy” per sensor; powers hindsight, insight & next-hour foresight.",
        ),
        (
            "Sensor locations",
            "Joins counts to the map; inactive/moved sensors must not mislead users.",
        ),
        (
            "Landmarks",
            "Epic 2 US 2.1: candidate refuges (parks/libraries)—not proof they are quiet.",
        ),
        (
            "OpenStreetMap",
            "Counts are points, not paths. OSM enables walkable sensory-weighted routes.",
        ),
    ]
    y = Inches(1.05)
    for title, desc in items:
        add_text_box(
            slide, Inches(0.4), y, Inches(12.4), Inches(0.28), f"• {title}", 12, True, INDIGO
        )
        add_text_box(slide, Inches(0.65), y + Inches(0.28), Inches(12.0), Inches(0.45), desc, 10)
        y += Inches(0.82)


def slide_customer_value(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GREY)
    add_title_bar(slide, "How each source creates customer value")

    w = Inches(4.05)
    add_bullet_box(
        slide,
        Inches(0.35),
        Inches(1.05),
        w,
        Inches(3.35),
        "Hindsight — What happened?",
        [
            "Recurring peak periods",
            "Quieter weekday/time combinations",
            "Seasonal & location differences",
            "Benefit: plan a historically quieter window",
        ],
    )
    add_bullet_box(
        slide,
        Inches(4.65),
        Inches(1.05),
        w,
        Inches(3.35),
        "Insight — What is happening?",
        [
            "Current relative crowd level",
            "Monitored lower-activity areas",
            "Nearby candidate refuges",
            "Benefit: avoid an unexpected crowd corridor",
        ],
    )
    add_bullet_box(
        slide,
        Inches(8.95),
        Inches(1.05),
        w,
        Inches(3.35),
        "Foresight — What may happen?",
        [
            "Next-hour crowd estimate",
            "Confidence / uncertainty",
            "Early rerouting prompt",
            "Benefit: adjust before conditions worsen",
        ],
    )

    banner = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.35),
        Inches(4.55),
        Inches(12.6),
        Inches(0.65),
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(0xFF, 0xED, 0xE5)
    banner.line.color.rgb = ORANGE
    tf = banner.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = (
        "Pedestrian volume is a sensory-load proxy—not a guarantee that a route or place is “safe”."
    )
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def slide_pipeline(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GREY)
    add_title_bar(slide, "Planned data pipeline and quality controls")

    steps = [
        "Official APIs\n+ OSM snapshot",
        "Immutable raw\nsnapshot + metadata",
        "Validate,\ndeduplicate,\nquality-flag",
        "Normalised\nPostgreSQL",
        "Baseline, forecast,\nroute weighting",
        "Application API\nmap & alerts",
    ]
    x = Inches(0.25)
    box_w = Inches(1.95)
    for i, label in enumerate(steps):
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(1.15), box_w, Inches(0.95)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_CYAN
        shape.line.color.rgb = CYAN
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(8)
        p.font.color.rgb = SLATE
        p.alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        if i < len(steps) - 1:
            slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                x + box_w + Inches(0.02),
                Inches(1.45),
                Inches(0.18),
                Inches(0.25),
            ).fill.solid()
        x += box_w + Inches(0.22)

    add_bullet_box(
        slide,
        Inches(0.35),
        Inches(2.35),
        Inches(6.15),
        Inches(2.55),
        "Ingestion controls",
        [
            "Record source URL, retrieval time and licence",
            "Validate IDs, timestamps, coordinates, non-negative counts",
            "Upsert on (sensor_id, observed_at)",
            "Quarantine invalid records—never silently discard",
        ],
    )
    add_bullet_box(
        slide,
        Inches(6.85),
        Inches(2.35),
        Inches(6.15),
        Inches(2.55),
        "Interpretation controls",
        [
            "Distinguish zero movement from sensor/API failure",
            "Preserve sensor relocation history and active periods",
            "Display freshness, coverage and prediction confidence",
            "Version raw snapshots, transforms and model outputs",
        ],
    )


def slide_erd(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GREY)
    add_title_bar(slide, "Planned ERD: source data, geography and derived insight")

    entities = [
        ("SENSOR", "sensor_id, status, notes", False),
        ("SENSOR_LOCATION_HISTORY", "location_id, sensor_id, lat/lon, valid_from/to", False),
        ("MINUTE_COUNT", "sensor_id, observed_at, totals, quality_status", False),
        ("HOURLY_COUNT", "sensor_id, observed_hour, total, quality_status", False),
        ("REFUGE_CANDIDATE", "place_id, category, name, coords, validation", False),
        ("ROAD_EDGE", "edge_id, from/to node, distance, walkability", False),
        ("CROWD_BASELINE", "sensor_id, weekday, hour, percentiles", True),
        ("CROWD_PREDICTION", "sensor_id, target_hour, level, confidence", True),
    ]
    positions = [
        (1.0, 1.1),
        (5.0, 1.1),
        (9.0, 1.1),
        (1.0, 2.55),
        (5.0, 2.55),
        (9.0, 2.55),
        (2.5, 4.05),
        (7.5, 4.05),
    ]
    for (name, fields, derived), (lx, ty) in zip(entities, positions):
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(lx),
            Inches(ty),
            Inches(3.1),
            Inches(1.15),
        )
        fill = RGBColor(0xFF, 0xF7, 0xED) if derived else LIGHT_GREY
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = ORANGE if derived else NAVY
        tf = shape.text_frame
        tf.clear()
        t = tf.paragraphs[0]
        t.text = name
        t.font.bold = True
        t.font.size = Pt(9)
        t.font.color.rgb = ORANGE if derived else NAVY
        b = tf.add_paragraph()
        b.text = fields
        b.font.size = Pt(8)
        b.font.color.rgb = SLATE

    add_text_box(
        slide,
        Inches(0.35),
        Inches(5.35),
        Inches(12.6),
        Inches(0.55),
        "Grey = imported facts. Orange = versioned analytical outputs. Location history prevents wrong map positions after sensor moves.",
        font_size=10,
        color=SLATE,
    )


def slide_analytics(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GREY)
    add_title_bar(slide, "Selected analytical approach: explainable first")

    add_bullet_box(
        slide,
        Inches(0.35),
        Inches(1.05),
        Inches(6.15),
        Inches(2.85),
        "Crowd classification (sensor-relative)",
        [
            "Low: ≤ 40th percentile for that sensor, weekday & hour",
            "Moderate: 40th–75th percentile",
            "High: > 75th percentile",
            "Avoids one global threshold for different street widths",
        ],
    )
    add_bullet_box(
        slide,
        Inches(6.85),
        Inches(1.05),
        Inches(6.15),
        Inches(2.85),
        "Next-hour forecast",
        [
            "1. Historical median for sensor + weekday + hour",
            "2. Adjust using most recent observation",
            "3. Confidence from sample size & data freshness",
            "4. Validate with MAE and alert precision/recall",
        ],
    )

    add_text_box(
        slide,
        Inches(0.35),
        Inches(4.05),
        Inches(12.6),
        Inches(0.7),
        "AI decision: no opaque ML required for onboarding. Add complexity only if validation shows better forecasts without losing explainability.",
        font_size=11,
        bold=True,
        color=NAVY,
    )

    alert = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.35),
        Inches(4.85),
        Inches(12.6),
        Inches(0.95),
    )
    alert.fill.solid()
    alert.fill.fore_color.rgb = RGBColor(0xFF, 0xED, 0xE5)
    alert.line.color.rgb = ORANGE
    tf = alert.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = (
        "Responsible interpretation: say “lower observed pedestrian activity” and “candidate refuge”—"
        "not “sensory-safe”. Show stale-data warnings, coverage gaps and uncertainty."
    )
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE


def slide_risks(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GREY)
    add_title_bar(slide, "Data risks, mitigations and source links")

    risks = [
        (
            "CBD sensor coverage gaps",
            "Show monitored area only; do not extrapolate beyond defensible distance.",
        ),
        (
            "Missing readings ambiguous",
            "Check sensor status & neighbours; use unavailable/uncertain—not blind interpolation.",
        ),
        (
            "Events break patterns",
            "Lower forecast confidence; note factors outside current datasets.",
        ),
        (
            "Facility may not be quiet/open",
            "Treat landmarks as candidates until validated.",
        ),
        (
            "Stale or altered public data",
            "Schema validation, checksummed snapshots, ingestion monitoring.",
        ),
    ]
    tbl = slide.shapes.add_table(
        len(risks) + 1, 2, Inches(0.35), Inches(1.05), Inches(12.6), Inches(3.2)
    ).table
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(9.4)
    for ci, h in enumerate(["Material risk", "Planned mitigation"]):
        c = tbl.cell(0, ci)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = INDIGO
        for p in c.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(10)
            p.font.color.rgb = WHITE
    for ri, (r, m) in enumerate(risks, 1):
        tbl.cell(ri, 0).text = r
        tbl.cell(ri, 1).text = m
        for ci in (0, 1):
            for p in tbl.cell(ri, ci).text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = SLATE

    links = (
        "Official sources (Aug 2026):\n"
        "• Minute counts: data.melbourne.vic.gov.au/.../pedestrian-counting-system-past-hour-counts-per-minute\n"
        "• Hourly counts: data.melbourne.vic.gov.au/.../pedestrian-counting-system-monthly-counts-per-hour\n"
        "• Sensor locations: data.melbourne.vic.gov.au/.../pedestrian-counting-system-sensor-locations\n"
        "• Landmarks: data.melbourne.vic.gov.au/.../landmarks-and-places-of-interest-...\n"
        "• OpenStreetMap (ODbL) • CC BY 4.0"
    )
    add_text_box(slide, Inches(0.35), Inches(4.35), Inches(12.6), Inches(1.5), links, 8, color=CYAN)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_datasets(prs)
    slide_why_chosen(prs)
    slide_customer_value(prs)
    slide_pipeline(prs)
    slide_erd(prs)
    slide_analytics(prs)
    slide_risks(prs)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
